# -*- coding: utf-8 -*-
"""Erzeugt eine Excel-Lead-Liste: Hamburger Firmen, die aktuell Grafiker/
Mediengestalter suchen (Basis fuer BENCON.group-Ansprache)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ---------------------------------------------------------------- Styles
BENCON_BLUE = "1F3864"
ACCENT = "2E5496"
HEADER_FILL = PatternFill("solid", fgColor=BENCON_BLUE)
ALT_FILL = PatternFill("solid", fgColor="EAF0FA")
FLAG_FILL = PatternFill("solid", fgColor="FCE8D5")
white_bold = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
title_font = Font(name="Calibri", bold=True, color=BENCON_BLUE, size=16)
sub_font = Font(name="Calibri", italic=True, color="555555", size=10)
cell_font = Font(name="Calibri", size=10)
link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
thin = Side(style="thin", color="BFC9DA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")
wrap_top_center = Alignment(wrap_text=True, vertical="top", horizontal="center")

# ================================================================ Sheet 1
ws = wb.active
ws.title = "Leads Hamburg"

ws["A1"] = "Lead-Liste: Hamburger Firmen mit offenen Stellen fuer Grafiker / Mediengestalter"
ws["A1"].font = title_font
ws["A2"] = ("Zweck: Aktive Ansprache durch BENCON.group (www.bencon.group) mit dem Angebot eines "
            "flexiblen Rahmenvertrags | Recherche-Stand: 04.07.2026 | Quellen: StepStone, Indeed, "
            "Arbeitsagentur, dasauge, Firmen-Impressum")
ws["A2"].font = sub_font

headers = [
    ("Nr.", 5),
    ("Firma / Arbeitgeber", 26),
    ("Gesuchte Position (lt. Anzeige)", 30),
    ("Was gesucht wird / Aufgaben", 40),
    ("Qualifikation / Anforderungen", 38),
    ("Angaben zum Arbeitgeber", 34),
    ("Datum der Anzeige", 15),
    ("Indiz: findet offenbar keinen", 26),
    ("Ansprechpartner", 20),
    ("E-Mail", 27),
    ("Telefon", 20),
    ("Adresse", 30),
    ("Website", 26),
    ("Quelle (Portal)", 22),
]

HEADER_ROW = 4
for col, (name, width) in enumerate(headers, start=1):
    c = ws.cell(row=HEADER_ROW, column=col, value=name)
    c.fill = HEADER_FILL
    c.font = white_bold
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c.border = border
    ws.column_dimensions[get_column_letter(col)].width = width

# ---------------------------------------------------------------- Data
# Felder: firma, position, aufgaben, quali, arbeitgeber, datum, indiz,
#         ansprechpartner, email, tel, adresse, website, quelle
leads = [
    ["your pac GmbH & Co. KG",
     "Mediengestalter (m/w/d) - Schwerpunkt Repro/Pre-Print",
     "Unterstuetzung des Repro-Teams bei Erstellung und finaler Aufbereitung von Druckdaten fuer hochwertige (flexible) Verpackungsprojekte.",
     "Ausbildung Mediengestalter o. vergleichbar; Erfahrung im Pre-Print; Adobe Creative Cloud; PackZ oder ArtPro.",
     "Digitaldruck / Verpackungsproduktion. Sitz: Bahrenfeld. GF Oliver Renz (HRA 128889).",
     "laufend 2026",
     "Nischen-Repro-Profil, dauerhaft ausgeschrieben",
     "GF Oliver Renz (kein Name in Anzeige)",
     "info@yourpac.de",
     "+49 40 298 101 78-0",
     "Stahltwiete 21a, 22761 Hamburg",
     "www.yourpac.de",
     "Indeed / Firmen-Website"],

    ["Backshop Tiefkuehl GmbH",
     "Mediengestalter (m/w/d)",
     "Kreative Gestaltung/Umsetzung innovativer Designs fuer Print & Digital; Broschueren & Kataloge; Bilddatenmanagement; Pflege der Produktqualitaet auf der Website.",
     "Abgeschlossene Ausbildung/Studium Mediengestalter o. vergleichbar; sicher in Adobe Creative Suite (InDesign, Photoshop, Illustrator).",
     "Grosshandel TK-Backwaren (Harry-Brot-Gruppe). Sitz: Bahrenfeld. GF Christian Sommer (HRB 105357).",
     "akt. ~Juni 2026 (vor 16 Tagen)",
     "mehrfach parallel (StepStone/LinkedIn/OMR), ab sofort",
     "kein Name genannt",
     "info@backshop-tk.de",
     "040 523876-0",
     "Leverkusenstr. 54, 22761 Hamburg-Bahrenfeld",
     "www.backshop-tk.de",
     "karriere-hamburg.de / StepStone"],

    ["Backstage Make-up GmbH",
     "Mediengestalter / Grafiker / Kommunikationsdesigner (m/w/d)",
     "Grafik/Design fuer Marke & Sortiment (Kosmetik); Bewerbung mit Arbeitsproben, Gehaltsvorstellung und Eintrittstermin erbeten.",
     "Ausbildung/Studium im Bereich Grafik/Mediengestaltung/Kommunikationsdesign; Portfolio.",
     "Kosmetik / Make-up (Handel & Studios). Sitz: Eppendorf/Eimsbuettel. GF Abraham Simonian (HRB 87752).",
     "laufend 2026",
     "breit gefasstes Profil (3 Berufe in 1 Anzeige)",
     "Verena (lt. Anzeige)",
     "info@backstage-make-up.de",
     "040 43910816",
     "Eppendorfer Weg 64, 20259 Hamburg",
     "www.backstage-makeup.de",
     "trabajo.org / dasauge"],

    ["YOOtheme GmbH",
     "Grafik Designer (m/w/d)",
     "Gestaltung fuer Software-/Web-Produkte (Themes, UIkit); UI-/Marketing-Design.",
     "Sehr gute Kenntnisse in Photoshop, Illustrator, XD und Figma.",
     "Softwarehersteller (Joomla/WordPress-Themes, Open-Source UIkit), ~700 m2 Buero HafenCity. GF Niels Nuebel, Stefan Wendhausen (HRB 108112).",
     "laufend 2026",
     "spezielles UI/UX-Toolset gefordert",
     "kein Name genannt",
     "info@yootheme.com",
     "+49 40 60943030-0",
     "Hongkongstrasse 10a, 20457 Hamburg",
     "www.yootheme.com",
     "dasauge / Firmen-Website"],

    ["A.S. Mueller Sofortdruck GmbH",
     "Kommunikationsdesigner / Grafik / Reinzeichnung (m/w/d), Voll-/Teilzeit",
     "Grafik, Reinzeichnung und Gestaltung fuer Druckauftraege.",
     "Kenntnisse in Grafikdesign / Reinzeichnung; Druckvorstufe.",
     "Druckerei (>50 Jahre). Sitz: Niendorf. Inhaber Hamid Reza Zaroofchi (HRB 35916).",
     "laufend 2026",
     "kleine Druckerei, schwer zu besetzen",
     "Inh. H. R. Zaroofchi",
     "Impressum pruefen (asm-druck.de)",
     "040 384043 / 0171 7043437",
     "Papenreye 17, 22453 Hamburg-Niendorf",
     "www.asm-druck.de",
     "dasauge / Firmen-Website"],

    ["Trade Con GmbH",
     "Grafiker:in / Art Director:in (w/m/d) sowie Mediengestalter:in (m/w/d) FR Printmedien",
     "Gestaltung/Reinzeichnung fuer Produkt-/Verpackungs- und Handelsmarken (Non-Food, Lizenzartikel).",
     "Grafik/Art Direction bzw. Mediengestaltung Printmedien; sicher in Adobe.",
     "Non-Food-Produktentwicklung & -beschaffung (>25 Jahre, Lizenzen UEFA CL/DFB). Sitz: Neustadt.",
     "ab 01.08. (mehrf. ausgeschr.)",
     "2 offene Stellen parallel (Grafik + Media)",
     "kein Name genannt",
     "kontakt@tradecon.net",
     "+49 40 2576642-10",
     "ABC-Strasse 21, 20354 Hamburg",
     "www.tradecon.net",
     "arbeitsagentur / dasauge"],

    ["Druckcenter Hamburg Pietro Piazza",
     "Mediengestalter Digital und Print (m/w/d) - Gestaltung u. Technik / Printmedien",
     "Gestaltung und technische Umsetzung von Print- und Digitalprodukten in inhabergefuehrter Druckerei.",
     "Ausbildung Mediengestalter Digital und Print o. vergleichbar; Druckvorstufe.",
     "Inhabergefuehrte Druckerei Eimsbuettel. Inhaber Pietro Piazza (USt DE216925712).",
     "laufend 2026",
     "Fachkraft UND Azubi gleichzeitig gesucht -> starkes Indiz",
     "Pietro Piazza (Inhaber)",
     "info@druckcenter-hamburg.de",
     "040 41090725",
     "Osterstrasse 22, 20259 Hamburg",
     "www.druckcenter-hamburg.de",
     "arbeitsagentur / LinkedIn / job.fish"],

    ["elbgraphen GmbH (Werbeagentur)",
     "Junior Art Director (w/m/d) Print",
     "Konzeption/Gestaltung von B2B-Kampagnen (Messe, Print, Internet); Reinzeichnung.",
     "Ausbildung/Studium Grafik/Kommunikationsdesign; Adobe; erste Berufserfahrung (Junior).",
     "B2B-Werbeagentur. Sitz: Altona/Ottensen. GF Stefan Klug (HRB 80513).",
     "Start ab 01.06.2026",
     "seit Fruehjahr offen, Junior-Level",
     "GF Stefan Klug",
     "info@elbgraphen.de",
     "040 38629200",
     "Van-der-Smissen-Str. 4, 22767 Hamburg",
     "www.elbgraphen.de",
     "elbgraphen.de / dasauge"],

    ["bung.art GmbH & Co. KG",
     "(Junior) Graphic Designer (m/w/d)",
     "Gestaltung/Umsetzung crossmedialer Kampagnen von Social Media ueber Print bis Web.",
     "Ausbildung/Studium Grafik-/Kommunikationsdesign; Adobe CC; crossmediales Arbeiten.",
     "Inhabergefuehrte Branding-/Kreativagentur (Hamburg & Berlin). GF Philip Bungart.",
     "laufend 2026",
     "wachsende Agentur, laufend im Recruiting",
     "GF Philip Bungart",
     "hallo@bung.art",
     "+49 40 3750 3870",
     "Katharinenstr. 4, 20457 Hamburg",
     "www.bung.art",
     "bung.art/jobs"],

    ["EINSATZ Creative Production GmbH & Co. KG",
     "Mediengestalter Digital und Print oder Reinzeichner Packaging (w/m/d)",
     "Erstellung von Masteradaptionen und Reinzeichnungen fuer Verpackung und Kommunikationsmittel nach Kundenvorgaben.",
     "Ausbildung Mediengestalter/Reinzeichner; sicher in Adobe; Packaging-Erfahrung von Vorteil.",
     "Creative Production (Packaging, Kataloge, Magazine, Artwork, Bildbearbeitung, 3D). Sitz: St. Pauli. GF Juliane Schmid (HRA 83479).",
     "laufend 2026",
     "spezialisiertes Packaging-Reinzeichner-Profil",
     "kein Name genannt",
     "einsatz@einsatz.de",
     "+49 40 376655",
     "Pinnasberg 47, 20359 Hamburg",
     "www.einsatz.de",
     "StepStone / Indeed"],

    ["brandport GmbH",
     "Mediengestalter / Reinzeichner Packaging (w/m/d)",
     "Packaging, Repro und Artwork; Reinzeichnung fuer Verpackung & POS.",
     "Erfahrung mit Esko ArtPro oder Esko DeskPack; Illustrator und InDesign.",
     "Repro/Prepress fuer Packaging & POS (Fusion Hirte Prepress + Hanse Reprozentrum, M+M-Gruppe). Sitz: Ottensen. GF Simon Frenz, Dieter Dolezal (HRB 32615).",
     "28.05.2026",
     "sehr spezielles Esko-Toolset -> schwer besetzbar",
     "kein Name genannt",
     "info@brandport.com",
     "+49 40 547770-0",
     "Stresemannstr. 375, Haus 9, 22761 Hamburg",
     "www.brandport.com",
     "Indeed / StepStone"],

    ["Cordes Consulting GmbH",
     "Mediengestalter / Mediendesign (m/w/d)",
     "Gestaltung von Kampagnen/Creatives im Social-Media-/Performance-Marketing; Agentur baut Mediendesign-Team aus.",
     "Ausbildung/Studium Mediengestaltung/Grafik; Adobe; Social-/Digital-Affinitaet.",
     "Digitalagentur (Recruiting & Marketing fuer Mittelstand/Handwerk, >300 Kunden), expandiert. GF David Cordes (HRB 172847).",
     "laufend 2025/2026",
     "expandiert, stellt lfd. in Mediendesign ein",
     "GF David Cordes",
     "kontakt@cordes-consulting.de",
     "040 239685190",
     "Esplanade 41, 20354 Hamburg",
     "www.cordes-consulting.de",
     "Indeed / Firmen-Website"],
]

r = HEADER_ROW + 1
for i, row in enumerate(leads, start=1):
    values = [i] + row
    for col, val in enumerate(values, start=1):
        c = ws.cell(row=r, column=col, value=val)
        c.font = cell_font
        c.border = border
        c.alignment = wrap_top_center if col in (1, 7) else wrap_top
        if i % 2 == 0:
            c.fill = ALT_FILL
    # Indiz-Spalte (H = 8) hervorheben
    ws.cell(row=r, column=8).fill = FLAG_FILL
    # E-Mail als klickbaren mailto-Link, wenn echte Adresse
    email = ws.cell(row=r, column=10)
    if email.value and "@" in str(email.value) and "pruefen" not in str(email.value):
        email.hyperlink = "mailto:" + str(email.value)
        email.font = link_font
    # Website-Link
    web = ws.cell(row=r, column=13)
    if web.value and web.value.startswith("www."):
        web.hyperlink = "https://" + web.value
        web.font = link_font
    r += 1

ws.freeze_panes = "B5"
ws.auto_filter.ref = f"A{HEADER_ROW}:N{r-1}"
ws.row_dimensions[HEADER_ROW].height = 34
for rr in range(HEADER_ROW + 1, r):
    ws.row_dimensions[rr].height = 78

# ================================================================ Sheet 2
ws2 = wb.create_sheet("Hinweise & Methodik")
ws2.column_dimensions["A"].width = 110
notes = [
    ("Hinweise zur Lead-Liste", title_font),
    ("", None),
    ("Ziel", Font(bold=True, color=BENCON_BLUE, size=12)),
    ("Hamburger Firmen identifizieren, die aktuell aktiv Grafiker / Mediengestalter in Jobportalen suchen und diese "
     "Stellen offenbar nur schwer besetzen koennen - als Basis fuer die aktive Ansprache durch die BENCON.group "
     "(www.bencon.group) mit dem Angebot eines flexiblen Rahmenvertrags.", cell_font),
    ("", None),
    ("Wie wurde recherchiert?", Font(bold=True, color=BENCON_BLUE, size=12)),
    ("- Auswertung offener Stellenanzeigen (Stand 04.07.2026) auf StepStone, Indeed, Arbeitsagentur, dasauge, "
     "karriere-hamburg.de u.a. mit Suchbegriffen Grafiker / Mediengestalter / Reinzeichner / Kommunikationsdesigner "
     "+ Hamburg.", cell_font),
    ("- Kontaktdaten (Adresse, E-Mail, Telefon, Geschaeftsfuehrung) stammen aus dem jeweiligen Firmen-Impressum bzw. "
     "aus Handelsregister-/Firmenauskunftsdiensten (Creditreform, Northdata, hamburg.de-Branchenbuch). Es wurden "
     "KEINE Daten erfunden.", cell_font),
    ("", None),
    ("Wichtig zur Spalte 'Indiz: findet offenbar keinen'", Font(bold=True, color="B45309", size=12)),
    ("Ob eine Firma 'niemanden findet', laesst sich von aussen nicht zu 100 % belegen. Als Indizien wurden gewertet: "
     "Anzeige laeuft laenger / wird mehrfach parallel auf mehreren Portalen ausgeschrieben; 'ab sofort' / 'zum "
     "naechstmoeglichen Zeitpunkt'; parallel Fachkraft UND Ausbildungsplatz; sehr spezielle Tool-Anforderungen "
     "(z.B. Esko ArtPro, PackZ). Vor der Ansprache empfiehlt sich eine kurze Einzelpruefung der jeweils aktuellen Anzeige.", cell_font),
    ("", None),
    ("Datenschutz / Ansprache", Font(bold=True, color=BENCON_BLUE, size=12)),
    ("Aufgefuehrt sind ausschliesslich geschaeftliche Kontaktdaten aus oeffentlichen Impressen/Anzeigen. Namentliche "
     "Ansprechpartner sind nur eingetragen, wenn sie in der Anzeige/im Impressum genannt waren. Bei der Kaltansprache "
     "(insb. E-Mail/Telefon) bitte die Vorgaben der DSGVO und des UWG (B2B-Werbung) beachten.", cell_font),
    ("", None),
    ("Empfohlene naechste Schritte", Font(bold=True, color=BENCON_BLUE, size=12)),
    ("1. Aktuelle Anzeige je Firma final pruefen (Link/Portal in letzter Spalte).\n"
     "2. Kurzes, individualisiertes Anschreiben je Firma: flexibler BENCON-Rahmenvertrag als Loesung fuer schwer "
     "besetzbare Grafik-/Mediengestalter-Stellen.\n"
     "3. Fehlende Ansprechpartner ggf. via Telefon/LinkedIn ermitteln.", cell_font),
]
rr = 1
for text, font in notes:
    c = ws2.cell(row=rr, column=1, value=text)
    if font:
        c.font = font
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.row_dimensions[rr].height = 30 if text and len(text) > 90 else 16
    rr += 1

out = "/home/user/prompt-eng-interactive-tutorial/Leads_Hamburg_Grafiker_Mediengestalter.xlsx"
wb.save(out)
print("gespeichert:", out)
print("Firmen:", len(leads))
