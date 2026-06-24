"""Excel-Ein-/Ausgabe und Datenmodell.

Liest die Personen-/Rollen-Tabelle ein und stellt eine Beispiel-Vorlage bereit.
Spaltennamen werden tolerant erkannt (deutsch/englisch, Groß/Klein egal).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import openpyxl

# Akzeptierte Spalten-Überschriften (normalisiert: lower, ohne Leer-/Sonderzeichen)
_ALIASES = {
    "name": "name",
    "officername": "name",
    "person": "name",
    "bedeutung": "role",
    "rolle": "role",
    "role": "role",
    "function": "role",
    "funktion": "role",
    "team": "team",
    "kategorie": "team",
    "category": "team",
    "officerid": "officer_id",
    "id": "officer_id",
    "nummer": "officer_id",
    "number": "officer_id",
    "site": "site",
    "standort": "site",
    "qr": "qr",
    "qrcode": "qr",
    "qrcontent": "qr",
    "qrdata": "qr",
    "url": "qr",
    "farbe": "color",
    "color": "color",
    "stickerfarbe": "color",
    "stickercolor": "color",
    "leiste": "color",
    "leistenfarbe": "color",
}

ID_START = 74401  # Start-Nummer wie im Original-Booklet


@dataclass
class Person:
    name: str = ""
    role: str = ""          # "Bedeutung" / Rolle, z.B. "ENTRANCE EXIT A"
    team: str = ""          # Kategorie, z.B. "CONTROL", "CBRN", "CSI", "DVI"
    officer_id: str = ""    # z.B. "74401"
    site: str = ""
    qr: str = ""            # expliziter QR-Inhalt (optional)
    color: str = ""         # Farbe der Seitenleiste, z.B. "#045B74" (optional)
    _row: int = field(default=0, repr=False)

    def qr_payload(self, base_url: Optional[str]) -> str:
        """Bestimme den QR-Inhalt mit folgender Priorität:

        1. explizite Spalte ``QR`` in der Tabelle,
        2. ``base_url`` + Officer-ID (falls --qr-base-url gesetzt),
        3. Fallback ``"Officer ID <id>"`` wie im Original.
        """
        if self.qr.strip():
            return self.qr.strip()
        if base_url:
            sep = "" if base_url.endswith(("/", "=", "?", "&")) else "/"
            return f"{base_url}{sep}{self.officer_id}"
        return f"Officer ID {self.officer_id}"


def _norm(key: str) -> str:
    return "".join(ch for ch in str(key).lower() if ch.isalnum())


def read_people(path: str | Path, id_start: int = ID_START) -> list[Person]:
    """Lies die Excel-Tabelle und liefere eine Liste von :class:`Person`.

    Fehlende Officer-IDs werden fortlaufend ab ``id_start`` vergeben.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Die Excel-Tabelle ist leer.")

    header, *data_rows = rows
    col_map: dict[int, str] = {}
    for idx, cell in enumerate(header):
        if cell is None:
            continue
        target = _ALIASES.get(_norm(cell))
        if target:
            col_map[idx] = target
    if "name" not in col_map.values() and "role" not in col_map.values():
        raise ValueError(
            "Es wurde weder eine 'Name'- noch eine 'Bedeutung'-Spalte erkannt. "
            f"Gefundene Überschriften: {[h for h in header if h]}"
        )

    people: list[Person] = []
    next_id = id_start
    for r, raw in enumerate(data_rows, start=2):
        values = {col_map[i]: ("" if raw[i] is None else str(raw[i]).strip())
                  for i in col_map if i < len(raw)}
        # Komplett leere Zeilen überspringen.
        if not any(values.get(k) for k in ("name", "role", "team", "officer_id")):
            continue
        p = Person(
            name=values.get("name", ""),
            role=values.get("role", ""),
            team=values.get("team", ""),
            officer_id=values.get("officer_id", ""),
            site=values.get("site", ""),
            qr=values.get("qr", ""),
            color=values.get("color", ""),
            _row=r,
        )
        if not p.officer_id:
            p.officer_id = str(next_id)
            next_id += 1
        else:
            # Auto-Zähler hinter eine evtl. numerische ID setzen.
            try:
                next_id = max(next_id, int(p.officer_id) + 1)
            except ValueError:
                pass
        people.append(p)
    if not people:
        raise ValueError("Keine Datenzeilen in der Tabelle gefunden.")
    return people


def write_example(path: str | Path) -> Path:
    """Schreibe eine Beispiel-/Blanko-Vorlage als .xlsx und gib den Pfad zurück."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Teams"
    headers = ["Name", "Bedeutung", "Team", "OfficerID", "Site", "QR", "Farbe"]
    ws.append(headers)
    examples = [
        ["Max Mustermann", "ENTRANCE EXIT A", "CONTROL", "74401", "", "", ""],
        ["Erika Beispiel", "LOGISTIC", "LOGISTIC", "74403", "", "", ""],
        ["John Doe", "CSI COMMANDER", "CSI", "74405", "", "https://etaf.example/o/74405", "#045B74"],
        ["Jane Roe", "DVI TEAMLEADER", "DVI", "74431", "", "", ""],
    ]
    for row in examples:
        ws.append(row)
    # Spaltenbreiten hübsch machen.
    for col, width in zip("ABCDEFG", (22, 26, 14, 12, 10, 34, 12)):
        ws.column_dimensions[col].width = width
    # Kommentar/Legende auf einem zweiten Blatt.
    info = wb.create_sheet("Hinweise")
    for line in [
        ["Spalte", "Bedeutung", "Pflicht?"],
        ["Name", "Name der Person (wird ins Namensfeld gedruckt)", "empfohlen"],
        ["Bedeutung", "Rolle/Funktion, z.B. 'ENTRANCE EXIT A'", "ja"],
        ["Team", "Kategorie, z.B. CONTROL / CBRN / CSI / DVI", "optional"],
        ["OfficerID", "Nummer; leer = automatisch ab 74401", "optional"],
        ["Site", "Standort/Einsatzort", "optional"],
        ["QR", "Eigener QR-Inhalt (URL/Text); leer = automatisch", "optional"],
        ["Farbe", "Aufkleber-/Leistenfarbe rechts, z.B. #045B74; leer = Originalfarbe", "optional"],
    ]:
        info.append(line)
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 50
    info.column_dimensions["C"].width = 12
    path = Path(path)
    wb.save(path)
    return path


def write_roster_template(path: str | Path, roster: list[dict]) -> Path:
    """Schreibe eine **vorbefüllte** Maske: je Rollen-Seite eine Zeile mit
    Officer-ID, Rolle und Team (aus dem Booklet ausgelesen). Die Spalten
    ``Name`` und ``QR`` bleiben leer und werden vom Nutzer ergänzt."""
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Officer"
    headers = ["Name", "Bedeutung", "Team", "OfficerID", "Site", "QR", "Farbe"]
    ws.append(headers)
    title = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="045B74")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = title
        cell.fill = fill
    for entry in roster:
        color = entry.get("color", "") or ""
        ws.append(["", entry.get("role", ""), entry.get("team", ""),
                   entry.get("officer_id", ""), "", "", color])
        # Die Farb-Zelle zur Sicht-Kontrolle mit der erkannten Farbe hinterlegen.
        hexv = color.lstrip("#")
        if len(hexv) == 6:
            cell = ws.cell(row=ws.max_row, column=7)
            cell.fill = PatternFill("solid", fgColor=hexv.upper())
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFG", (24, 28, 14, 12, 12, 34, 12)):
        ws.column_dimensions[col].width = width

    info = wb.create_sheet("Hinweise")
    for line in [
        ["So füllst du die Maske aus"],
        [""],
        ["• Jede Zeile ist eine Rollen-Seite des Booklets (Reihenfolge wie im Original)."],
        ["• Trage in Spalte 'Name' den Namen der Person ein – mehr ist nicht nötig."],
        ["• 'Bedeutung'/'Team'/'OfficerID' sind schon vorbefüllt (zur Orientierung)."],
        ["• 'QR' optional: eigener QR-Inhalt (URL/Text); leer = automatisch 'Officer ID <id>'."],
        ["• 'Farbe' = Farbe des Aufklebers/der Seitenleiste rechts (z.B. #045B74),"],
        ["   schon aus dem Booklet ausgelesen. Leer lassen = Originalfarbe bleibt."],
        ["   Für eine zusätzliche Person die passende Farbe eintragen (#RRGGBB)."],
        ["• Zeilen ohne Namen kannst du löschen, wenn die Rolle nicht besetzt ist."],
    ]:
        info.append(line)
    info.column_dimensions["A"].width = 80
    path = Path(path)
    wb.save(path)
    return path
